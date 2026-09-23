#!/usr/bin/env python3
"""Refresh tools/pawn-desk-prices.xlsx against what the app actually prices.

The master sheet is how the counter gets its OWN numbers into the desk
without a single API call: fill the shaded YOUR VALUE (or YOUR LOW / YOUR
HIGH) column, save that one sheet as CSV, and drop it on the Setup tab. A
number entered that way stands in front of every published one, counts as
good data and never goes stale.

The workbook was built by hand once and drifted the moment the catalog and
the harvest moved: rows for things the desk no longer prices, and nothing
at all for the ones it learned since. A drifted key does not fail loudly -
the importer just lists it as skipped, and the counter's afternoon of
typing goes nowhere.

This UPDATES the workbook rather than rebuilding it. The layout is good and
was not mine: the Kind/Market/Sells columns, the live Difference formula,
the footer counts and the Read me all stay exactly as they are. Only the
data rows are rewritten, and the Market note on a row that already existed
is carried across rather than guessed at again.

    node tools/dump-prices.mjs > dump.json
    python3 tools/build-price-sheet.py dump.json

The heading row is not decoration - app.js matches on those words. A sheet
carrying YOUR LOW and YOUR HIGH is read as the models sheet and one with a
single YOUR VALUE as everything else, which is what stops the model row
keyed "a1" being swallowed by the appliance keyed "a1".
"""
import json, sys, os, collections
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(HERE, "tools", "pawn-desk-prices.xlsx")
src  = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "tools", "price-dump.json")
d    = json.load(open(src))

FILLME = "FFF6DC"
thin   = Side(style="thin", color="D5DDD8")
box    = Border(left=thin, right=thin, top=thin, bottom=thin)
body   = Font(size=10, name="Calibri", color="1A1A1A")

wb = load_workbook(XLSX)
SHEET = {c["label"]: c for c in d["cats"]}
book_by_cat = collections.defaultdict(list)
for b in d["book"]:
    book_by_cat[b["cat"]].append(b)

def sheet_for(name):
    """The workbook's tab names are shortened by hand - "Appliances" for
       "Appliances & household", "Cards & collectibles" for "Cards, coins &
       collectibles" - so neither name reliably starts with the other. The
       first word does the matching, and it is unique across all eleven."""
    if name in SHEET: return SHEET[name]
    def first(s): return s.split(" ")[0].strip(",&").lower()
    for label, c in SHEET.items():
        if first(label) == first(name):
            return c
    return None

def clear(ws, first, last, width):
    for r in range(first, last + 1):
        for c in range(1, width + 1):
            ws.cell(row=r, column=c).value = None

def style(ws, r, width, fill_cols):
    for c in range(1, width + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = box
        cell.font = body
        if c in fill_cols:
            cell.fill = PatternFill("solid", fgColor=FILLME)

report = []
for name in wb.sheetnames:
    if name in ("Read me", "Models"): continue
    cat = sheet_for(name)
    if not cat:
        report.append(f"  {name}: no category matched - left alone"); continue
    ws = wb[name]
    old_rows = list(ws.iter_rows(min_row=3, values_only=True))
    was = {str(r[1]): r for r in old_rows if r[1]}
    market = collections.Counter(r[7] for r in old_rows if r[1] and r[7])
    default_market = market.most_common(1)[0][0] if market else "ships"

    want = [("List item", i["id"], i["name"], i["value"], i["liq"]) for i in cat["items"]]
    want += [("Price book", b["name"], b["name"], b["val"], b["liq"]) for b in book_by_cat[cat["id"]]]

    clear(ws, 3, max(len(old_rows) + 6, len(want) + 6) + 3, 9)
    r = 3
    for kind, key, item, val, liq in want:
        prev = was.get(str(key))
        ws.cell(row=r, column=1, value=kind)
        ws.cell(row=r, column=2, value=key)
        ws.cell(row=r, column=3, value=item)
        ws.cell(row=r, column=4, value=val)
        ws.cell(row=r, column=6, value=f'=IF(E{r}="","",E{r}-D{r})')
        ws.cell(row=r, column=7, value=liq)
        ws.cell(row=r, column=8, value=(prev[7] if prev and prev[7] else default_market))
        style(ws, r, 9, {5, 9})
        r += 1
    last = r - 1
    ws.cell(row=last + 2, column=3, value="Rows on this sheet")
    ws.cell(row=last + 2, column=4, value=f"=COUNTA(C3:C{last})")
    ws.cell(row=last + 3, column=3, value="Ones you changed")
    ws.cell(row=last + 3, column=4, value=f"=COUNT(E3:E{last})")
    ws.cell(row=last + 5, column=1,
            value="Edit the yellow columns only. Leave a row blank and it stays exactly as it is.")
    added = [k for _, k, _, _, _ in want if str(k) not in was]
    gone  = [k for k in was if k not in {str(k2) for _, k2, _, _, _ in want}]
    report.append(f"  {name:<24} {len(want):>4} rows   +{len(added)} new  -{len(gone)} dropped")

# ---- Models --------------------------------------------------------------
ws = wb["Models"]
old = list(ws.iter_rows(min_row=3, values_only=True))
was = {str(r[0]): r for r in old if r[0]}
clear(ws, 3, max(len(old), len(d["models"])) + 8, 11)
r = 3
for m in sorted(d["models"], key=lambda x: (x["ref"], x["name"])):
    ws.cell(row=r, column=1, value=m["id"])
    ws.cell(row=r, column=2, value=m["ref"])
    ws.cell(row=r, column=3, value=m["name"])
    ws.cell(row=r, column=4, value=m["lo"])
    ws.cell(row=r, column=5, value=m["hi"])
    ws.cell(row=r, column=8, value=m["conf"])
    ws.cell(row=r, column=9, value=m["date"])
    ws.cell(row=r, column=10, value=m.get("src", ""))
    ws.cell(row=r, column=11, value=m.get("note", ""))
    style(ws, r, 11, {6, 7})
    r += 1
last = r - 1
ws.cell(row=last + 2, column=3, value="Model rows")
ws.cell(row=last + 2, column=4, value=f"=COUNTA(C3:C{last})")
added = [m["id"] for m in d["models"] if m["id"] not in was]
gone  = [k for k in was if k not in {m["id"] for m in d["models"]}]
report.append(f"  {'Models':<24} {len(d['models']):>4} rows   +{len(added)} new  -{len(gone)} dropped")

wb.save(XLSX)
n = sum(len(c["items"]) for c in d["cats"]) + len(d["book"]) + len(d["models"])
print("refreshed " + XLSX)
print("\n".join(report))
print(f"\n  {n} priced things in all "
      f"({sum(len(c['items']) for c in d['cats'])} items, {len(d['book'])} price book, {len(d['models'])} models)")
